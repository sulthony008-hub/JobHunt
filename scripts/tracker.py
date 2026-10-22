"""Helper tracker lamaran (Excel) untuk agen JobHunter.

Contoh:
    python tracker.py list
    python tracker.py list --status Terkirim
    python tracker.py add --perusahaan "PT Contoh" --posisi "Data Analyst" \
        --lokasi "Jakarta (Hybrid)" --sumber JobStreet --email "hrd@contoh.co.id"
    python tracker.py update --no 1 --status Terkirim --subjek "Lamaran ..." \
        --message-id "<abc@mail.gmail.com>"
    python tracker.py log --aksi "Cari lowongan" --detail "5 hasil baru"
"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(os.environ.get("JOBHUNTER_ROOT", Path(__file__).resolve().parent.parent))
TRACKER = Path(os.environ.get("JOBHUNTER_TRACKER", ROOT / "lamaran.xlsx"))
SHEET = "Lamaran"
FIRST_ROW = 2
LAST_ROW = 301

STATUSES = [
    "Ditemukan",
    "Disiapkan",
    "Terkirim",
    "Follow-up",
    "Screening",
    "Interview",
    "Tes",
    "Offer",
    "Ditolak",
    "Tidak Ada Balasan",
    "Ditarik",
]

COLS = {
    "no": 1,
    "tanggal_ditemukan": 2,
    "perusahaan": 3,
    "posisi": 4,
    "lokasi": 5,
    "sumber": 6,
    "url": 7,
    "email": 8,
    "status": 9,
    "tanggal_kirim": 10,
    "subjek": 11,
    "message_id": 12,
    "tahap": 13,
    "deadline": 14,
    "catatan": 15,
    "update_terakhir": 16,
}


def today() -> str:
    return date.today().isoformat()


def load():
    if not TRACKER.exists():
        sys.exit(f"Tracker tidak ditemukan: {TRACKER}")
    wb = load_workbook(TRACKER)
    return wb, wb[SHEET]


def row_number(ws, row: int):
    """Nomor lamaran efektif (kolom A berisi formula)."""
    value = ws.cell(row=row, column=COLS["no"]).value
    if isinstance(value, (int, float)):
        return int(value)
    if isinstance(value, str) and value.startswith("="):
        return row - 1
    return None


def find_row(ws, no: int) -> int:
    for row in range(FIRST_ROW, LAST_ROW + 1):
        if row_number(ws, row) == no:
            return row
    sys.exit(f"Lamaran nomor {no} tidak ditemukan.")


def first_empty_row(ws) -> int:
    for row in range(FIRST_ROW, LAST_ROW + 1):
        if not ws.cell(row=row, column=COLS["perusahaan"]).value:
            return row
    sys.exit("Tracker penuh (maksimal 300 baris).")


def set_cell(ws, row: int, field: str, value):
    if value is None:
        return
    ws.cell(row=row, column=COLS[field], value=value)


def append_log(ws_log, aksi: str, detail: str, oleh: str = "hermes"):
    row = 2
    while ws_log.cell(row=row, column=1).value:
        row += 1
    ws_log.cell(row=row, column=1, value=f"{today()} {datetime.now():%H:%M}")
    ws_log.cell(row=row, column=2, value=aksi)
    ws_log.cell(row=row, column=3, value=detail)
    ws_log.cell(row=row, column=4, value=oleh)


def cmd_add(args):
    wb, ws = load()
    row = first_empty_row(ws)
    values = {
        "tanggal_ditemukan": args.tanggal_ditemukan or today(),
        "perusahaan": args.perusahaan,
        "posisi": args.posisi,
        "lokasi": args.lokasi,
        "sumber": args.sumber,
        "url": args.url,
        "email": args.email,
        "status": args.status,
        "tanggal_kirim": args.tanggal_kirim,
        "subjek": args.subjek,
        "message_id": args.message_id,
        "tahap": args.tahap,
        "deadline": args.deadline,
        "catatan": args.catatan,
        "update_terakhir": today(),
    }
    for field, value in values.items():
        set_cell(ws, row, field, value)
    append_log(wb["Log"], "Tambah lamaran", f"{args.perusahaan} - {args.posisi} ({args.status})")
    wb.save(TRACKER)
    print(f"OK: lamaran #{row - 1} ditambahkan -> {args.perusahaan} - {args.posisi}")


def cmd_update(args):
    wb, ws = load()
    row = find_row(ws, args.no)
    changed = []
    for field in [
        "perusahaan",
        "posisi",
        "lokasi",
        "sumber",
        "url",
        "email",
        "status",
        "tanggal_kirim",
        "subjek",
        "message_id",
        "tahap",
        "deadline",
        "catatan",
    ]:
        value = getattr(args, field, None)
        if value is not None:
            set_cell(ws, row, field, value)
            changed.append(field)

    if args.status == "Terkirim" and not ws.cell(row=row, column=COLS["tanggal_kirim"]).value:
        set_cell(ws, row, "tanggal_kirim", today())
        changed.append("tanggal_kirim")

    set_cell(ws, row, "update_terakhir", today())
    append_log(wb["Log"], "Update lamaran", f"#{args.no}: {', '.join(changed) or 'tanpa perubahan'}")
    wb.save(TRACKER)
    print(f"OK: lamaran #{args.no} diperbarui ({', '.join(changed) or 'tanpa perubahan'})")


def cmd_list(args):
    _, ws = load()
    rows = []
    for row in range(FIRST_ROW, LAST_ROW + 1):
        no = row_number(ws, row)
        if no is None:
            continue
        if not ws.cell(row=row, column=COLS["perusahaan"]).value and not ws.cell(
            row=row, column=COLS["status"]
        ).value:
            continue
        status = ws.cell(row=row, column=COLS["status"]).value or ""
        if args.status and status != args.status:
            continue
        rows.append(
            {
                "no": no,
                "status": status,
                "perusahaan": ws.cell(row=row, column=COLS["perusahaan"]).value or "",
                "posisi": ws.cell(row=row, column=COLS["posisi"]).value or "",
                "lokasi": ws.cell(row=row, column=COLS["lokasi"]).value or "",
                "sumber": ws.cell(row=row, column=COLS["sumber"]).value or "",
                "email": ws.cell(row=row, column=COLS["email"]).value or "",
                "tanggal_kirim": ws.cell(row=row, column=COLS["tanggal_kirim"]).value or "",
                "tahap": ws.cell(row=row, column=COLS["tahap"]).value or "",
                "deadline": ws.cell(row=row, column=COLS["deadline"]).value or "",
                "catatan": ws.cell(row=row, column=COLS["catatan"]).value or "",
            }
        )

    if args.json:
        print(json.dumps(rows, ensure_ascii=False, indent=2))
        return

    if not rows:
        print("Belum ada lamaran.")
        return

    print(f"{'No':>3}  {'Status':<18} {'Perusahaan':<24} {'Posisi':<28} {'Kirim':<12} Tahap")
    print("-" * 110)
    for item in rows:
        print(
            f"{item['no']:>3}  {item['status']:<18} {item['perusahaan'][:23]:<24} "
            f"{item['posisi'][:27]:<28} {str(item['tanggal_kirim'])[:10]:<12} {item['tahap']}"
        )


def cmd_log(args):
    wb, _ = load()
    append_log(wb["Log"], args.aksi, args.detail, args.oleh)
    wb.save(TRACKER)
    print("OK: log ditambahkan.")


def build_parser():
    p = argparse.ArgumentParser(description="Tracker lamaran JobHunter")
    sub = p.add_subparsers(dest="command", required=True)

    a = sub.add_parser("add", help="Tambah lamaran baru")
    a.add_argument("--perusahaan", required=True)
    a.add_argument("--posisi", required=True)
    a.add_argument("--lokasi")
    a.add_argument("--sumber")
    a.add_argument("--url")
    a.add_argument("--email")
    a.add_argument("--status", default="Ditemukan", choices=STATUSES)
    a.add_argument("--tanggal-ditemukan", dest="tanggal_ditemukan")
    a.add_argument("--tanggal-kirim", dest="tanggal_kirim")
    a.add_argument("--subjek")
    a.add_argument("--message-id", dest="message_id")
    a.add_argument("--tahap")
    a.add_argument("--deadline")
    a.add_argument("--catatan")
    a.set_defaults(func=cmd_add)

    u = sub.add_parser("update", help="Update lamaran")
    u.add_argument("--no", type=int, required=True)
    u.add_argument("--perusahaan")
    u.add_argument("--posisi")
    u.add_argument("--lokasi")
    u.add_argument("--sumber")
    u.add_argument("--url")
    u.add_argument("--email")
    u.add_argument("--status", choices=STATUSES)
    u.add_argument("--tanggal-kirim", dest="tanggal_kirim")
    u.add_argument("--subjek")
    u.add_argument("--message-id", dest="message_id")
    u.add_argument("--tahap")
    u.add_argument("--deadline")
    u.add_argument("--catatan")
    u.set_defaults(func=cmd_update)

    l = sub.add_parser("list", help="Lihat lamaran")
    l.add_argument("--status", choices=STATUSES)
    l.add_argument("--json", action="store_true")
    l.set_defaults(func=cmd_list)

    g = sub.add_parser("log", help="Tambah catatan log")
    g.add_argument("--aksi", required=True)
    g.add_argument("--detail", required=True)
    g.add_argument("--oleh", default="hermes")
    g.set_defaults(func=cmd_log)
    return p


if __name__ == "__main__":
    args = build_parser().parse_args()
    args.func(args)

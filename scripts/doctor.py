"""Cek kesiapan JobHunter tanpa mengirim email.

    python scripts/doctor.py
    python scripts/doctor.py --online

--online mencoba koneksi IMAP (butuh config + app password).
"""

from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(os.environ.get("JOBHUNTER_ROOT", Path(__file__).resolve().parent.parent))
PROFIL = ROOT / "profil.yaml"
SUMBER = ROOT / "sumber.yaml"
TRACKER = ROOT / "lamaran.xlsx"
CV = ROOT / "cv" / "CV.pdf"
DRAFTS = ROOT / "drafts"
TEMPLATES = ROOT / "templates"
DASHBOARD = ROOT / "dashboard"
HIMALAYA = Path(os.environ.get("HIMALAYA_BIN", shutil.which("himalaya") or r"D:\Hermes\bin\himalaya.exe"))
HERMES_JOBS = Path(r"D:\Hermes\cron\jobs.json")
CONFIG = Path.home() / ".config" / "himalaya" / "config.toml"
PASSWORD = Path.home() / ".config" / "himalaya" / "app-password.txt"

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
problems: list[str] = []


def ok(label: str, detail: str = "") -> None:
    print(f"[OK] {label}" + (f" â€” {detail}" if detail else ""))


def warn(label: str, detail: str = "") -> None:
    print(f"[!!] {label}" + (f" â€” {detail}" if detail else ""))
    problems.append(label)


def info(label: str, detail: str = "") -> None:
    print(f"[--] {label}" + (f" â€” {detail}" if detail else ""))


def check_profile() -> None:
    try:
        import yaml
    except ImportError:
        warn("PyYAML tidak terpasang", "pip install pyyaml")
        return
    if not PROFIL.exists():
        warn("profil.yaml tidak ada")
        return
    data = yaml.safe_load(PROFIL.read_text(encoding="utf-8")) or {}
    required = ["nama_lengkap", "email_agen", "no_hp", "kota_domisili"]
    missing = [key for key in required if not str(data.get(key) or "").strip()]
    if missing:
        warn("Profil belum lengkap", ", ".join(missing))
    else:
        ok("Profil terisi", data["nama_lengkap"])

    email = str(data.get("email_agen") or "")
    if email and not EMAIL_RE.match(email):
        warn("email_agen tidak valid", email)

    kriteria = data.get("kriteria") or {}
    posisi = kriteria.get("posisi") or []
    lokasi = kriteria.get("lokasi") or []
    if not posisi:
        warn("Kriteria posisi kosong")
    else:
        ok("Kriteria posisi", ", ".join(posisi))
    if not lokasi:
        warn("Kriteria lokasi kosong")
    else:
        ok("Kriteria lokasi", ", ".join(lokasi))


def check_cv() -> None:
    if not CV.exists():
        warn("CV belum ada", str(CV))
        return
    size = CV.stat().st_size
    if size < 10_000:
        warn("CV terlalu kecil / mungkin bukan PDF", f"{size} bytes")
    else:
        ok("CV tersedia", f"{size // 1024} KB")


def check_tracker() -> None:
    if not TRACKER.exists():
        warn("Tracker tidak ada", str(TRACKER))
        return
    try:
        from openpyxl import load_workbook
    except ImportError:
        warn("openpyxl tidak terpasang", "pip install openpyxl")
        return
    wb = load_workbook(TRACKER, read_only=True)
    if "Lamaran" not in wb.sheetnames:
        warn("Sheet 'Lamaran' tidak ditemukan", ", ".join(wb.sheetnames))
        return
    ws = wb["Lamaran"]
    rows = 0
    statuses: dict[str, int] = {}
    for row in ws.iter_rows(min_row=2, max_col=9, values_only=True):
        if row[2] or row[8]:
            rows += 1
            status = str(row[8] or "(kosong)")
            statuses[status] = statuses.get(status, 0) + 1
    wb.close()
    if rows:
        ringkas = ", ".join(f"{key}={value}" for key, value in sorted(statuses.items()))
        ok("Tracker terbaca", f"{rows} baris ({ringkas})")
    else:
        ok("Tracker siap", "belum ada lamaran")


def check_email(online: bool) -> None:
    if not HIMALAYA.exists():
        warn("himalaya tidak ditemukan", str(HIMALAYA))
    else:
        try:
            result = subprocess.run(
                [str(HIMALAYA), "--version"], capture_output=True, text=True, timeout=15
            )
            version = (result.stdout or result.stderr).strip().splitlines()[0]
            ok("himalaya", version)
        except (OSError, subprocess.TimeoutExpired) as exc:
            warn("himalaya gagal dijalankan", str(exc))

    if not CONFIG.exists():
        warn("Config email belum ada", str(CONFIG))
        return
    ok("Config email ada", str(CONFIG))

    if not PASSWORD.exists():
        warn("App password belum ada", str(PASSWORD))
        return
    password = PASSWORD.read_text(encoding="utf-8").strip()
    if len(password) != 16:
        warn("App password bukan 16 karakter", f"{len(password)} karakter")
    else:
        ok("App password tersedia", "16 karakter")

    if not online:
        info("Koneksi IMAP tidak diuji", "jalankan dengan --online")
        return
    try:
        result = subprocess.run(
            [str(HIMALAYA), "envelope", "list", "-a", "gmail", "-f", "INBOX", "-s", "1", "-o", "json"],
            capture_output=True,
            text=True,
            timeout=45,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        warn("Koneksi IMAP gagal", str(exc))
        return
    if result.returncode == 0:
        ok("Koneksi IMAP", "INBOX terbaca")
    else:
        detail = (result.stderr or result.stdout).strip()
        warn("Koneksi IMAP gagal", detail[:300])
        if "Invalid credentials" in detail or "AUTHENTICATIONFAILED" in detail:
            info(
                "Petunjuk",
                "Gmail menolak kredensial â€” pakai App Password 16 karakter, bukan password akun.",
            )


def check_files() -> None:
    for path in [
        TEMPLATES / "email_lamaran.txt",
        TEMPLATES / "follow_up.txt",
        ROOT / "scripts" / "tracker.py",
    ]:
        if path.exists():
            ok(path.name, str(path))
        else:
            warn("File hilang", str(path))
    if DRAFTS.exists():
        count = len(list(DRAFTS.glob("*.txt")))
        info("Draft", f"{count} file di {DRAFTS}")
    else:
        warn("Folder drafts tidak ada", str(DRAFTS))


def check_sources() -> None:
    if not SUMBER.exists():
        info("sumber.yaml belum ada", "dashboard akan membuat daftar default")
        return
    try:
        import yaml

        data = yaml.safe_load(SUMBER.read_text(encoding="utf-8")) or {}
    except Exception:
        warn("sumber.yaml tidak terbaca")
        return
    sources = data.get("sources") or []
    active = [item for item in sources if isinstance(item, dict) and item.get("aktif", True)]
    ok("Sumber lowongan", f"{len(sources)} total Â· {len(active)} aktif")


def check_dashboard() -> None:
    for path in [
        DASHBOARD / "app.py",
        DASHBOARD / "templates" / "index.html",
        DASHBOARD / "static" / "style.css",
        DASHBOARD / "static" / "app.js",
        ROOT / "Buka Dashboard.bat",
    ]:
        if not path.exists():
            warn("File dashboard hilang", str(path))
            return
    try:
        import flask  # noqa: F401
        from importlib.metadata import version

        ok("Dashboard", f"Flask {version('flask')} Â· http://127.0.0.1:8770")
    except ImportError:
        warn("Flask tidak terpasang", "pip install flask")


def check_cron() -> None:
    if not HERMES_JOBS.exists():
        info("Cron Hermes tidak terbaca", str(HERMES_JOBS))
        return
    try:
        jobs = json.loads(HERMES_JOBS.read_text(encoding="utf-8")).get("jobs", [])
    except (OSError, json.JSONDecodeError):
        info("Cron Hermes tidak terbaca", str(HERMES_JOBS))
        return
    hunt = [job for job in jobs if "jobhunt" in str(job.get("name", "")).lower()]
    if not hunt:
        info("Cron JobHunter", "belum dibuat")
        return
    for job in hunt:
        state = "aktif" if job.get("enabled") else "paused"
        ok(f"Cron {job.get('name')}", f"{state} Â· {job.get('schedule_display')}")


def main() -> int:
    online = "--online" in sys.argv[1:]
    print("=== Doctor JobHunter ===\n")
    check_profile()
    check_cv()
    check_email(online)
    check_tracker()
    check_files()
    check_dashboard()
    check_sources()
    check_cron()

    print()
    if problems:
        print(f"SELESAI: {len(problems)} hal perlu dibereskan.")
        for item in problems:
            print(f"  - {item}")
        return 1
    print("SELESAI: semua siap.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

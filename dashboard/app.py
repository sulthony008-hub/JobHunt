"""Dashboard lokal JobHunter.

Menjalankan:
    python dashboard/app.py
Lalu buka http://127.0.0.1:8770

Dashboard ini hanya mendengarkan di localhost. Tidak ada data yang dikirim
keluar selain permintaan tes koneksi IMAP ke Gmail.
"""

from __future__ import annotations

import importlib.util
import json
import math
import os
import re
import shutil
import subprocess
import sys
import threading
import tomllib
from datetime import date, datetime
from pathlib import Path

import yaml
from flask import Flask, jsonify, render_template, request, send_file
from openpyxl import load_workbook

ROOT = Path(os.environ.get("JOBHUNTER_ROOT", Path(__file__).resolve().parent.parent))
SCRIPTS = ROOT / "scripts"
PROFIL = ROOT / "profil.yaml"
SUMBER = ROOT / "sumber.yaml"
CV = ROOT / "cv" / "CV.pdf"
CV_VERSIONS = ROOT / "cv" / "versions"
TEMPLATES = ROOT / "templates"
HIMALAYA = Path(os.environ.get("HIMALAYA_BIN", shutil.which("himalaya") or r"D:\Hermes\bin\himalaya.exe"))
VERIFIED = ROOT / ".email-verified.json"

sys.path.insert(0, str(SCRIPTS))
import tracker  # noqa: E402

TRACKER = tracker.TRACKER
STATUSES = tracker.STATUSES
COLS = tracker.COLS
FIRST_ROW = tracker.FIRST_ROW
LAST_ROW = tracker.LAST_ROW

_spec = importlib.util.spec_from_file_location("jhsetup", str(SCRIPTS / "setup.py"))
jhsetup = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(jhsetup)

CONFIG = jhsetup.CONFIG
PASSWORD = jhsetup.PASSWORD

app = Flask(__name__, static_folder="static", template_folder="templates")
app.config["MAX_CONTENT_LENGTH"] = 12 * 1024 * 1024
_lock = threading.Lock()

DEFAULT_SOURCES = [
    {
        "nama": "LinkedIn",
        "url": "https://www.linkedin.com/jobs/",
        "tipe": "Professional network",
        "catatan": "Cari lowongan; jangan auto-apply karena melanggar ToS.",
        "aktif": True,
    },
    {
        "nama": "JobStreet",
        "url": "https://www.jobstreet.co.id/",
        "tipe": "Job board",
        "catatan": "Lowongan Indonesia, banyak posisi entry-mid.",
        "aktif": True,
    },
    {
        "nama": "Glints",
        "url": "https://glints.com/id/opportunities/jobs/explore",
        "tipe": "Job board",
        "catatan": "",
        "aktif": True,
    },
    {
        "nama": "Kalibrr",
        "url": "https://www.kalibrr.com/",
        "tipe": "Job board",
        "catatan": "",
        "aktif": True,
    },
    {
        "nama": "Dealls",
        "url": "https://dealls.com/",
        "tipe": "Job board",
        "catatan": "",
        "aktif": True,
    },
    {
        "nama": "Indeed",
        "url": "https://id.indeed.com/",
        "tipe": "Job board",
        "catatan": "",
        "aktif": True,
    },
    {
        "nama": "Career page perusahaan",
        "url": "",
        "tipe": "Company career page",
        "catatan": "Cek langsung halaman karier perusahaan target.",
        "aktif": True,
    },
]


# ------------------------------------------------------------------ helpers


def now_iso() -> str:
    return datetime.now().isoformat(timespec="seconds")


def today_str() -> str:
    return date.today().isoformat()


def fmt_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and math.isfinite(value) and value.is_integer():
        return int(value)
    return value


def parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def read_yaml(path: Path) -> dict:
    if not path.exists():
        return {}
    try:
        data = yaml.safe_load(path.read_text(encoding="utf-8"))
    except (OSError, yaml.YAMLError):
        return {}
    return data if isinstance(data, dict) else {}


def write_yaml(path: Path, data: dict, header: str = "") -> None:
    body = yaml.safe_dump(data, allow_unicode=True, sort_keys=False, default_flow_style=False)
    path.write_text(header + body, encoding="utf-8")


def save_profile(data: dict) -> None:
    header = (
        "# Profil pelamar & kriteria lowongan.\n"
        "# Dibuat/diperbarui oleh dashboard JobHunter. Boleh diedit manual.\n\n"
    )
    write_yaml(PROFIL, data, header)


def slugify(text: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return slug or "sumber"


def source_list() -> list[dict]:
    data = read_yaml(SUMBER)
    sources = data.get("sources")
    if not isinstance(sources, list) or not sources:
        sources = []
        profile = read_yaml(PROFIL)
        profile_names = [str(name).strip() for name in profile.get("sumber") or [] if str(name).strip()]
        by_name = {item["nama"].lower(): item for item in DEFAULT_SOURCES}
        for name in profile_names:
            template = by_name.get(name.lower())
            if template:
                sources.append(dict(template))
            else:
                sources.append(
                    {
                        "nama": name,
                        "url": "",
                        "tipe": "Lainnya",
                        "catatan": "",
                        "aktif": True,
                    }
                )
        for template in DEFAULT_SOURCES:
            if template["nama"].lower() not in {item["nama"].lower() for item in sources}:
                sources.append(dict(template))
        for index, item in enumerate(sources, start=1):
            item["id"] = slugify(item["nama"])
            item.setdefault("aktif", True)
            item.setdefault("catatan", "")
            item.setdefault("tipe", "Job board")
        write_yaml(
            SUMBER,
            {"sources": sources},
            "# Sumber lowongan JobHunter. Diedit lewat dashboard.\n\n",
        )
    else:
        cleaned = []
        seen = set()
        for index, item in enumerate(sources, start=1):
            if not isinstance(item, dict):
                continue
            item.setdefault("id", f"{slugify(str(item.get('nama', 'sumber')))}-{index}")
            while item["id"] in seen:
                item["id"] = f"{item['id']}-{index}"
            seen.add(item["id"])
            item.setdefault("nama", "Tanpa nama")
            item.setdefault("url", "")
            item.setdefault("tipe", "Job board")
            item.setdefault("catatan", "")
            item.setdefault("aktif", True)
            cleaned.append(item)
        sources = cleaned
        if not sources:
            for template in DEFAULT_SOURCES:
                sources.append(dict(template, id=slugify(template["nama"])))
            write_yaml(
                SUMBER,
                {"sources": sources},
                "# Sumber lowongan JobHunter. Diedit lewat dashboard.\n\n",
            )
    return sources


def save_sources(sources: list[dict]) -> None:
    write_yaml(
        SUMBER,
        {"sources": sources},
        "# Sumber lowongan JobHunter. Diedit lewat dashboard.\n\n",
    )
    profile = read_yaml(PROFIL)
    profile["sumber"] = [item["nama"] for item in sources if item.get("aktif", True)]
    save_profile(profile)


def cv_filename() -> str:
    profile = read_yaml(PROFIL)
    name = str(profile.get("nama_lengkap") or "Pelamar").strip()
    safe = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "Pelamar"
    return f"CV_{safe}.pdf"


def cv_meta() -> dict:
    if not CV.exists():
        return {"exists": False, "download_name": cv_filename()}
    stat = CV.stat()
    return {
        "exists": True,
        "size": stat.st_size,
        "size_label": f"{stat.st_size / 1024:.0f} KB",
        "modified": datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        "url": f"/api/cv/file?v={int(stat.st_mtime)}",
        "download_name": cv_filename(),
    }


def email_verified(email: str) -> bool:
    if not email or not VERIFIED.exists():
        return False
    try:
        data = json.loads(VERIFIED.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return False
    return str(data.get("email") or "").lower() == email.lower()


def email_status() -> dict:
    profile = read_yaml(PROFIL)
    config_email = ""
    if CONFIG.exists():
        try:
            config = tomllib.loads(CONFIG.read_text(encoding="utf-8"))
            config_email = (
                config.get("accounts", {}).get("gmail", {}).get("email", "") or ""
            )
        except (OSError, tomllib.TOMLDecodeError):
            config_email = ""
    password_ok = False
    if PASSWORD.exists():
        try:
            password_ok = len(PASSWORD.read_text(encoding="utf-8").strip()) == 16
        except OSError:
            password_ok = False
    email = config_email or str(profile.get("email_agen") or "")
    return {
        "email": email,
        "config_exists": CONFIG.exists(),
        "password_ok": password_ok,
        "himalaya_exists": HIMALAYA.exists(),
        "ready": CONFIG.exists() and password_ok and HIMALAYA.exists(),
        "verified": email_verified(email),
        "config_path": str(CONFIG),
        "password_path": str(PASSWORD),
    }


def read_lamaran() -> list[dict]:
    if not TRACKER.exists():
        return []
    wb = load_workbook(TRACKER, read_only=True, data_only=False)
    if "Lamaran" not in wb.sheetnames:
        wb.close()
        return []
    ws = wb["Lamaran"]
    items = []
    fields = list(COLS.items())
    max_col = max(COLS.values())
    rows = ws.iter_rows(
        min_row=FIRST_ROW,
        max_row=LAST_ROW,
        max_col=max_col,
        values_only=True,
    )
    for offset, row in enumerate(rows, start=FIRST_ROW):
        values = {
            field: fmt_value(row[col - 1] if col - 1 < len(row) else None)
            for field, col in fields
        }
        if not values.get("perusahaan") and not values.get("status"):
            continue
        values["no"] = offset - 1
        items.append(values)
    wb.close()
    return items


def load_tracker_for_write():
    if not TRACKER.exists():
        raise FileNotFoundError(f"Tracker tidak ditemukan: {TRACKER}")
    wb = load_workbook(TRACKER)
    if "Lamaran" not in wb.sheetnames:
        raise ValueError("Sheet 'Lamaran' tidak ditemukan.")
    return wb, wb["Lamaran"]


def save_tracker(wb) -> None:
    tmp = TRACKER.with_name(TRACKER.stem + ".dashboard-tmp.xlsx")
    wb.save(tmp)
    os.replace(tmp, TRACKER)


def append_log(wb, aksi: str, detail: str) -> None:
    if "Log" in wb.sheetnames:
        tracker.append_log(wb["Log"], aksi, detail, oleh="dashboard")


def profile_payload(data: dict) -> dict:
    kriteria = data.get("kriteria") or {}
    pengiriman = data.get("pengiriman") or {}
    notifikasi = data.get("notifikasi") or {}
    return {
        "nama_lengkap": data.get("nama_lengkap", ""),
        "email_agen": data.get("email_agen", ""),
        "no_hp": data.get("no_hp", ""),
        "kota_domisili": data.get("kota_domisili", ""),
        "linkedin": data.get("linkedin", ""),
        "portfolio": data.get("portfolio", ""),
        "kriteria": {
            "posisi": kriteria.get("posisi") or [],
            "lokasi": kriteria.get("lokasi") or [],
            "mode": kriteria.get("mode") or [],
            "gaji_min": kriteria.get("gaji_min"),
            "level": kriteria.get("level") or [],
            "industri": kriteria.get("industri") or [],
            "kata_kunci_wajib": kriteria.get("kata_kunci_wajib") or [],
            "kata_kunci_hindari": kriteria.get("kata_kunci_hindari") or [],
        },
        "pengiriman": {
            "mode": pengiriman.get("mode", "draft_approval"),
            "batas_per_hari": pengiriman.get("batas_per_hari", 10),
            "follow_up_hari": pengiriman.get("follow_up_hari", 7),
            "follow_up_maksimal": pengiriman.get("follow_up_maksimal", 2),
            "kirim_dari_email_agen": pengiriman.get("kirim_dari_email_agen", True),
        },
        "notifikasi": {"telegram": bool(notifikasi.get("telegram", False))},
        "tracker": data.get("tracker", str(TRACKER)),
    }


def profile_missing(data: dict) -> list[str]:
    missing = []
    for key, label in [
        ("nama_lengkap", "Nama lengkap"),
        ("email_agen", "Email agen"),
        ("no_hp", "No. HP"),
        ("kota_domisili", "Kota domisili"),
    ]:
        if not str(data.get(key) or "").strip():
            missing.append(label)
    kriteria = data.get("kriteria") or {}
    if not kriteria.get("posisi"):
        missing.append("Posisi yang dicari")
    if not kriteria.get("lokasi"):
        missing.append("Lokasi")
    return missing


def split_list(value) -> list[str]:
    if isinstance(value, list):
        items = value
    else:
        items = str(value or "").split(",")
    return [str(item).strip() for item in items if str(item).strip()]


# ------------------------------------------------------------------- routes


@app.get("/")
def index():
    return render_template("index.html")


@app.errorhandler(413)
def too_large(_error):
    return jsonify({"ok": False, "error": "File terlalu besar. Maksimal 10 MB."}), 413


@app.get("/api/summary")
def api_summary():
    items = read_lamaran()
    counts = {status: 0 for status in STATUSES}
    for item in items:
        status = str(item.get("status") or "")
        if status in counts:
            counts[status] += 1
    today = date.today()
    followup_due = []
    upcoming = []
    for item in items:
        status = str(item.get("status") or "")
        if status in ("Terkirim", "Follow-up"):
            sent = parse_date(item.get("tanggal_kirim"))
            if sent and (today - sent).days >= 7:
                followup_due.append(
                    {**item, "days": (today - sent).days}
                )
        if status in ("Interview", "Tes", "Offer"):
            upcoming.append(item)
    recent = sorted(
        items,
        key=lambda item: (str(item.get("update_terakhir") or ""), int(item.get("no") or 0)),
        reverse=True,
    )[:5]
    profile = read_yaml(PROFIL)
    sources = source_list()
    return jsonify(
        {
            "ok": True,
            "generated_at": now_iso(),
            "profile": {
                "nama_lengkap": profile.get("nama_lengkap", ""),
                "email_agen": profile.get("email_agen", ""),
                "missing": profile_missing(profile),
            },
            "statuses": STATUSES,
            "counts": counts,
            "total": len(items),
            "followup_due": followup_due,
            "upcoming": upcoming,
            "recent": recent,
            "cv": cv_meta(),
            "email": email_status(),
            "sources": {
                "total": len(sources),
                "active": len([item for item in sources if item.get("aktif", True)]),
            },
        }
    )


@app.get("/api/lamaran")
def api_lamaran():
    items = read_lamaran()
    query = str(request.args.get("q") or "").strip().lower()
    status = str(request.args.get("status") or "").strip()
    sumber = str(request.args.get("sumber") or "").strip().lower()
    if query:
        items = [
            item
            for item in items
            if query
            in " ".join(
                str(item.get(field) or "")
                for field in ("perusahaan", "posisi", "lokasi", "catatan", "email")
            ).lower()
        ]
    if status:
        items = [item for item in items if item.get("status") == status]
    if sumber:
        items = [item for item in items if str(item.get("sumber") or "").lower() == sumber]
    items.sort(key=lambda item: int(item.get("no") or 0), reverse=True)
    return jsonify({"ok": True, "items": items, "statuses": STATUSES})


@app.post("/api/lamaran")
def api_add_lamaran():
    payload = request.get_json(silent=True) or {}
    perusahaan = str(payload.get("perusahaan") or "").strip()
    posisi = str(payload.get("posisi") or "").strip()
    if not perusahaan or not posisi:
        return jsonify({"ok": False, "error": "Perusahaan dan posisi wajib diisi."}), 400
    status = str(payload.get("status") or "Ditemukan")
    if status not in STATUSES:
        return jsonify({"ok": False, "error": "Status tidak valid."}), 400
    with _lock:
        try:
            wb, ws = load_tracker_for_write()
        except (FileNotFoundError, ValueError) as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500
        try:
            row = tracker.first_empty_row(ws)
        except SystemExit:
            wb.close()
            return jsonify({"ok": False, "error": "Tracker penuh (maksimal 300 baris)."}), 400
        values = {
            "tanggal_ditemukan": str(payload.get("tanggal_ditemukan") or today_str()),
            "perusahaan": perusahaan,
            "posisi": posisi,
            "lokasi": str(payload.get("lokasi") or ""),
            "sumber": str(payload.get("sumber") or ""),
            "url": str(payload.get("url") or ""),
            "email": str(payload.get("email") or ""),
            "status": status,
            "tanggal_kirim": str(
                payload.get("tanggal_kirim") or (today_str() if status == "Terkirim" else "")
            ),
            "subjek": str(payload.get("subjek") or ""),
            "message_id": str(payload.get("message_id") or ""),
            "tahap": str(payload.get("tahap") or ""),
            "deadline": str(payload.get("deadline") or ""),
            "catatan": str(payload.get("catatan") or ""),
            "update_terakhir": today_str(),
        }
        for field, value in values.items():
            tracker.set_cell(ws, row, field, value or None)
        append_log(wb, "Tambah lamaran", f"{perusahaan} - {posisi} ({status})")
        try:
            save_tracker(wb)
        except OSError as exc:
            wb.close()
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "Tracker sedang dipakai. Tutup lamaran.xlsx di Excel lalu coba lagi.",
                        "detail": str(exc),
                    }
                ),
                409,
            )
        wb.close()
    number = row - 1
    return jsonify({"ok": True, "no": number, "message": f"Lamaran #{number} ditambahkan."})


@app.patch("/api/lamaran/<int:no>")
def api_update_lamaran(no: int):
    payload = request.get_json(silent=True) or {}
    allowed = [
        "perusahaan",
        "posisi",
        "lokasi",
        "sumber",
        "url",
        "email",
        "status",
        "tanggal_kirim",
        "subjek",
        "message_id",
        "tahap",
        "deadline",
        "catatan",
    ]
    if payload.get("status") and payload["status"] not in STATUSES:
        return jsonify({"ok": False, "error": "Status tidak valid."}), 400
    with _lock:
        try:
            wb, ws = load_tracker_for_write()
            row = tracker.find_row(ws, no)
        except (FileNotFoundError, ValueError) as exc:
            return jsonify({"ok": False, "error": str(exc)}), 500
        except SystemExit:
            return jsonify({"ok": False, "error": "Lamaran tidak ditemukan."}), 404
        changed = []
        for field in allowed:
            if field in payload:
                value = payload[field]
                if isinstance(value, str):
                    value = value.strip()
                tracker.set_cell(ws, row, field, value)
                changed.append(field)
        if payload.get("status") == "Terkirim" and not ws.cell(
            row=row, column=COLS["tanggal_kirim"]
        ).value:
            tracker.set_cell(ws, row, "tanggal_kirim", today_str())
            changed.append("tanggal_kirim")
        tracker.set_cell(ws, row, "update_terakhir", today_str())
        append_log(wb, "Update lamaran", f"#{no}: {', '.join(changed) or 'tanpa perubahan'}")
        try:
            save_tracker(wb)
        except OSError as exc:
            wb.close()
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "Tracker sedang dipakai. Tutup lamaran.xlsx di Excel lalu coba lagi.",
                        "detail": str(exc),
                    }
                ),
                409,
            )
        wb.close()
    return jsonify({"ok": True, "message": f"Lamaran #{no} diperbarui."})


@app.get("/api/profil")
def api_profil():
    data = read_yaml(PROFIL)
    return jsonify(
        {
            "ok": True,
            "profil": profile_payload(data),
            "missing": profile_missing(data),
            "email": email_status(),
            "tracker": str(TRACKER),
        }
    )


@app.put("/api/profil")
def api_save_profil():
    payload = request.get_json(silent=True) or {}
    data = read_yaml(PROFIL)
    for key in ("nama_lengkap", "email_agen", "no_hp", "kota_domisili", "linkedin", "portfolio"):
        if key in payload:
            data[key] = str(payload.get(key) or "").strip()
    email = str(data.get("email_agen") or "")
    if email and not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        return jsonify({"ok": False, "error": "Format email agen tidak valid."}), 400

    kriteria_in = payload.get("kriteria") or {}
    kriteria = data.get("kriteria") or {}
    for key in (
        "posisi",
        "lokasi",
        "mode",
        "level",
        "industri",
        "kata_kunci_wajib",
        "kata_kunci_hindari",
    ):
        if key in kriteria_in:
            kriteria[key] = split_list(kriteria_in.get(key))
    if "gaji_min" in kriteria_in:
        raw = str(kriteria_in.get("gaji_min") or "").replace(".", "").replace(",", "").strip()
        try:
            kriteria["gaji_min"] = int(raw) if raw else None
        except ValueError:
            return jsonify({"ok": False, "error": "Gaji minimum harus berupa angka."}), 400
    data["kriteria"] = kriteria

    pengiriman_in = payload.get("pengiriman") or {}
    pengiriman = data.get("pengiriman") or {}
    if "mode" in pengiriman_in:
        pengiriman["mode"] = str(pengiriman_in.get("mode") or "draft_approval")
    for key in ("batas_per_hari", "follow_up_hari", "follow_up_maksimal"):
        if key in pengiriman_in:
            try:
                pengiriman[key] = max(0, int(pengiriman_in.get(key) or 0))
            except (TypeError, ValueError):
                return jsonify({"ok": False, "error": f"{key} harus berupa angka."}), 400
    if "kirim_dari_email_agen" in pengiriman_in:
        pengiriman["kirim_dari_email_agen"] = bool(pengiriman_in.get("kirim_dari_email_agen"))
    data["pengiriman"] = pengiriman

    notifikasi = data.get("notifikasi") or {}
    notifikasi["telegram"] = bool((payload.get("notifikasi") or {}).get("telegram", notifikasi.get("telegram", False)))
    data["notifikasi"] = notifikasi
    data["cv_path"] = str(CV).replace("\\", "/")
    data["tracker"] = str(TRACKER).replace("\\", "/")

    save_profile(data)
    return jsonify(
        {
            "ok": True,
            "message": "Profil disimpan.",
            "profil": profile_payload(data),
            "missing": profile_missing(data),
        }
    )


@app.get("/api/cv")
def api_cv():
    return jsonify({"ok": True, "cv": cv_meta()})


@app.post("/api/cv")
def api_upload_cv():
    uploaded = request.files.get("cv")
    if uploaded is None:
        return jsonify({"ok": False, "error": "File CV tidak ditemukan."}), 400
    data = uploaded.read(10 * 1024 * 1024 + 1)
    if len(data) > 10 * 1024 * 1024:
        return jsonify({"ok": False, "error": "Ukuran CV maksimal 10 MB."}), 413
    if not data.startswith(b"%PDF"):
        return jsonify({"ok": False, "error": "File harus berupa PDF."}), 400
    CV.parent.mkdir(parents=True, exist_ok=True)
    with _lock:
        try:
            if CV.exists():
                CV_VERSIONS.mkdir(parents=True, exist_ok=True)
                stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
                shutil.copy2(CV, CV_VERSIONS / f"CV_{stamp}.pdf")
            tmp = CV.with_name("CV.uploading")
            tmp.write_bytes(data)
            os.replace(tmp, CV)
        except OSError as exc:
            return (
                jsonify(
                    {
                        "ok": False,
                        "error": "CV sedang dipakai atau tidak bisa ditulis. Tutup pratinjau CV lalu coba lagi.",
                        "detail": str(exc),
                    }
                ),
                409,
            )
    return jsonify({"ok": True, "message": "CV diperbarui.", "cv": cv_meta()})


@app.get("/api/cv/file")
def api_cv_file():
    if not CV.exists():
        return jsonify({"ok": False, "error": "CV belum ada."}), 404
    download = request.args.get("download") == "1"
    return send_file(
        CV,
        mimetype="application/pdf",
        as_attachment=download,
        download_name=cv_filename(),
        conditional=True,
    )


@app.get("/api/sumber")
def api_sumber():
    sources = source_list()
    return jsonify({"ok": True, "sources": sources})


@app.post("/api/sumber")
def api_add_sumber():
    payload = request.get_json(silent=True) or {}
    nama = str(payload.get("nama") or "").strip()
    if not nama:
        return jsonify({"ok": False, "error": "Nama sumber wajib diisi."}), 400
    url = str(payload.get("url") or "").strip()
    if url and not re.match(r"^https?://", url, re.I):
        return jsonify({"ok": False, "error": "URL harus diawali http:// atau https://"}), 400
    with _lock:
        sources = source_list()
        if any(str(item.get("nama", "")).lower() == nama.lower() for item in sources):
            return jsonify({"ok": False, "error": "Sumber dengan nama itu sudah ada."}), 400
        base = slugify(nama)
        source_id = base
        index = 2
        while any(item.get("id") == source_id for item in sources):
            source_id = f"{base}-{index}"
            index += 1
        sources.append(
            {
                "id": source_id,
                "nama": nama,
                "url": url,
                "tipe": str(payload.get("tipe") or "Job board"),
                "catatan": str(payload.get("catatan") or ""),
                "aktif": bool(payload.get("aktif", True)),
            }
        )
        save_sources(sources)
    return jsonify({"ok": True, "message": f"Sumber '{nama}' ditambahkan.", "sources": sources})


@app.put("/api/sumber/<source_id>")
def api_update_sumber(source_id: str):
    payload = request.get_json(silent=True) or {}
    with _lock:
        sources = source_list()
        target = next((item for item in sources if item.get("id") == source_id), None)
        if target is None:
            return jsonify({"ok": False, "error": "Sumber tidak ditemukan."}), 404
        if "nama" in payload:
            nama = str(payload.get("nama") or "").strip()
            if not nama:
                return jsonify({"ok": False, "error": "Nama sumber wajib diisi."}), 400
            if any(
                str(item.get("nama", "")).lower() == nama.lower()
                and item.get("id") != source_id
                for item in sources
            ):
                return jsonify({"ok": False, "error": "Nama sumber sudah dipakai."}), 400
            target["nama"] = nama
        if "url" in payload:
            url = str(payload.get("url") or "").strip()
            if url and not re.match(r"^https?://", url, re.I):
                return jsonify({"ok": False, "error": "URL harus diawali http:// atau https://"}), 400
            target["url"] = url
        if "tipe" in payload:
            target["tipe"] = str(payload.get("tipe") or "Job board")
        if "catatan" in payload:
            target["catatan"] = str(payload.get("catatan") or "")
        if "aktif" in payload:
            target["aktif"] = bool(payload.get("aktif"))
        save_sources(sources)
    return jsonify({"ok": True, "message": "Sumber diperbarui.", "sources": sources})


@app.delete("/api/sumber/<source_id>")
def api_delete_sumber(source_id: str):
    with _lock:
        sources = source_list()
        target = next((item for item in sources if item.get("id") == source_id), None)
        if target is None:
            return jsonify({"ok": False, "error": "Sumber tidak ditemukan."}), 404
        sources = [item for item in sources if item.get("id") != source_id]
        save_sources(sources)
    return jsonify({"ok": True, "message": f"Sumber '{target['nama']}' dihapus.", "sources": sources})


@app.get("/api/email")
def api_email():
    return jsonify({"ok": True, "email": email_status()})


@app.post("/api/email/test")
def api_email_test():
    if not HIMALAYA.exists():
        return jsonify({"ok": False, "error": f"himalaya tidak ditemukan: {HIMALAYA}"}), 400
    if not CONFIG.exists():
        return jsonify({"ok": False, "error": "Config email belum dibuat. Setel Gmail dulu."}), 400
    try:
        result = subprocess.run(
            [str(HIMALAYA), "envelope", "list", "-a", "gmail", "-f", "INBOX", "-s", "5", "-o", "json"],
            capture_output=True,
            text=True,
            timeout=45,
        )
    except (OSError, subprocess.TimeoutExpired) as exc:
        return jsonify({"ok": False, "error": f"Gagal menjalankan himalaya: {exc}"}), 500
    if result.returncode != 0:
        detail = (result.stderr or result.stdout).strip()
        if "Invalid credentials" in detail or "AUTHENTICATIONFAILED" in detail:
            message = (
                "Koneksi IMAP gagal: Gmail menolak kredensial. "
                "Pakai App Password 16 karakter (bukan password akun), "
                "pastikan 2FA aktif dan IMAP diaktifkan."
            )
        else:
            message = f"Koneksi IMAP gagal: {detail[:400]}"
        try:
            VERIFIED.unlink(missing_ok=True)
        except OSError:
            pass
        return (
            jsonify(
                {
                    "ok": False,
                    "error": message,
                    "detail": detail[:800],
                }
            ),
            400,
        )
    count = None
    try:
        parsed = json.loads(result.stdout or "[]")
        count = len(parsed) if isinstance(parsed, list) else None
    except json.JSONDecodeError:
        count = None
    try:
        jhsetup.sync_folder_aliases()
    except Exception:
        pass
    try:
        status = email_status()
        VERIFIED.write_text(
            json.dumps({"email": status.get("email", ""), "at": now_iso()}),
            encoding="utf-8",
        )
        status["verified"] = True
    except OSError:
        status = email_status()
    suffix = f" {count} email terbaru terbaca." if count is not None else " INBOX terbaca."
    return jsonify(
        {
            "ok": True,
            "message": "Koneksi IMAP berhasil." + suffix,
            "email": status,
        }
    )


@app.post("/api/email/setup")
def api_email_setup():
    payload = request.get_json(silent=True) or {}
    email = str(payload.get("email") or "").strip()
    password = re.sub(r"\s+", "", str(payload.get("password") or ""))
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", email):
        return jsonify({"ok": False, "error": "Format email tidak valid."}), 400
    if len(password) != 16:
        return jsonify({"ok": False, "error": "App Password harus 16 karakter (spasi otomatis dihapus)."}), 400
    with _lock:
        CONFIG.parent.mkdir(parents=True, exist_ok=True)
        PASSWORD.write_text(password + "\n", encoding="utf-8", newline="\n")
        jhsetup.restrict(PASSWORD)
        jhsetup.write_config(email)
        try:
            VERIFIED.unlink(missing_ok=True)
        except OSError:
            pass
        profile = read_yaml(PROFIL)
        profile["email_agen"] = email
        save_profile(profile)
    return jsonify(
        {
            "ok": True,
            "message": "Akun Gmail disimpan. Klik Tes koneksi untuk memverifikasi.",
            "email": email_status(),
        }
    )


@app.get("/api/health")
def api_health():
    return jsonify({"ok": True, "time": now_iso()})


if __name__ == "__main__":
    print("JobHunter dashboard: http://127.0.0.1:8770")
    app.run(host="127.0.0.1", port=8770, debug=False, use_reloader=False)
